from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.style import WD_STYLE_TYPE

OUT = 'AWS_Lambda_Scaling_Recommendation.docx'

def set_cell_shading(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = tcPr.find(qn('w:shd'))
    if shd is None:
        shd = OxmlElement('w:shd')
        tcPr.append(shd)
    shd.set(qn('w:fill'), fill)

def set_cell_border(cell, color='D9D9D9', size='6'):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    borders = tcPr.first_child_found_in('w:tcBorders')
    if borders is None:
        borders = OxmlElement('w:tcBorders')
        tcPr.append(borders)
    for edge in ('top','left','bottom','right','insideH','insideV'):
        tag = 'w:' + edge
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn('w:val'), 'single')
        element.set(qn('w:sz'), size)
        element.set(qn('w:space'), '0')
        element.set(qn('w:color'), color)

def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tcPr = tc.get_or_add_tcPr()
    tcMar = tcPr.first_child_found_in('w:tcMar')
    if tcMar is None:
        tcMar = OxmlElement('w:tcMar')
        tcPr.append(tcMar)
    for m, v in [('top',top),('start',start),('bottom',bottom),('end',end)]:
        node = tcMar.find(qn('w:' + m))
        if node is None:
            node = OxmlElement('w:' + m)
            tcMar.append(node)
        node.set(qn('w:w'), str(v)); node.set(qn('w:type'), 'dxa')

def set_repeat_table_header(row):
    trPr = row._tr.get_or_add_trPr()
    tblHeader = OxmlElement('w:tblHeader')
    tblHeader.set(qn('w:val'), 'true')
    trPr.append(tblHeader)

def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = 'Table Grid'
    hdr = table.rows[0]
    set_repeat_table_header(hdr)
    for i, text in enumerate(headers):
        c = hdr.cells[i]; c.text = text; c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
        set_cell_shading(c, '1F4E79'); set_cell_border(c); set_cell_margins(c)
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.bold = True; r.font.color.rgb = RGBColor(255,255,255); r.font.size = Pt(9)
    for ridx, row in enumerate(rows):
        cells = table.add_row().cells
        for i, text in enumerate(row):
            c = cells[i]; c.text = str(text); c.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            set_cell_border(c); set_cell_margins(c)
            if ridx % 2 == 1: set_cell_shading(c, 'F2F6FA')
            for p in c.paragraphs:
                p.paragraph_format.space_after = Pt(2)
                for r in p.runs: r.font.size = Pt(9)
    if widths:
        for row in table.rows:
            for i, w in enumerate(widths): row.cells[i].width = Inches(w)
    doc.add_paragraph().paragraph_format.space_after = Pt(1)
    return table

def add_bullet(doc, text, level=0):
    p = doc.add_paragraph(style='List Bullet' if level == 0 else 'List Bullet 2')
    p.paragraph_format.space_after = Pt(3)
    p.add_run(text)
    return p

doc = Document()
sec = doc.sections[0]
sec.top_margin = Inches(0.7); sec.bottom_margin = Inches(0.65)
sec.left_margin = Inches(0.8); sec.right_margin = Inches(0.8)

styles = doc.styles
styles['Normal'].font.name = 'Aptos'; styles['Normal'].font.size = Pt(10)
styles['Normal'].paragraph_format.space_after = Pt(6)
for name, size, color in [('Title', 25, '000000'), ('Heading 1', 16, '000000'), ('Heading 2', 12, '000000')]:
    st = styles[name]; st.font.name = 'Aptos Display' if name == 'Title' else 'Aptos'; st.font.size = Pt(size); st.font.bold = True; st.font.color.rgb = RGBColor.from_string(color)
    st.paragraph_format.space_before = Pt(12 if name != 'Title' else 0); st.paragraph_format.space_after = Pt(6)

title = doc.add_paragraph(style='Title'); title.alignment = WD_ALIGN_PARAGRAPH.LEFT
title.add_run('AWS Lambda Scaling Recommendation')
sub = doc.add_paragraph(); sub.paragraph_format.space_after = Pt(12)
r = sub.add_run('Serverless architecture for approximately 3,000 users with a target budget below 20 USD per month')
r.italic = True; r.font.size = Pt(11); r.font.color.rgb = RGBColor(80,80,80)

doc.add_heading('Recommendation', level=1)
doc.add_paragraph('Use AWS Lambda behind Amazon API Gateway HTTP API for the application API, while retaining Supabase for PostgreSQL and authentication. Store uploaded files in Amazon S3 using presigned URLs. Keep AWS IVS for video or live-streaming features, and run volume conversion as a separate asynchronous worker rather than inside the API Lambda.')
doc.add_paragraph('This approach is likely to remain below 20 USD per month for ordinary API traffic and a modest number of file operations. The budget does not automatically include high-volume IVS delivery, large S3 downloads, or compute-heavy volume conversion.')

doc.add_heading('Current application profile', level=1)
doc.add_paragraph('The repository contains a Node.js and Express API with many authenticated routes, a PostgreSQL connection through Supabase, Supabase client usage, Socket.IO, AWS IVS integrations, file uploads, and a separate Python volume-conversion workload. The API currently starts a long-running server on port 4004, so it needs a small adaptation before it can run on Lambda.')
add_table(doc, ['Area', 'Current implementation', 'Serverless implication'], [
    ('HTTP API', 'Express application with many route modules', 'Use one initial Lambda with an Express adapter, then split high-traffic routes if needed'),
    ('Database and auth', 'Supabase PostgreSQL and Supabase services', 'Keep Supabase; reuse clients and pooled connections outside the handler'),
    ('Realtime', 'Socket.IO and IVS-related routes', 'Use API Gateway WebSocket API, AppSync, or IVS realtime features'),
    ('Uploads', 'Multer memory upload path and storage integrations', 'Upload directly to S3 using short-lived presigned URLs'),
    ('Conversion', 'Python volume-to-NIfTI processing', 'Run asynchronously in a worker; do not block API requests'),
], widths=[1.2, 2.4, 3.2])

doc.add_heading('Target architecture', level=1)
doc.add_paragraph('The recommended request path is: frontend to API Gateway HTTP API, API Gateway to Lambda, and Lambda to Supabase. The frontend should request an S3 presigned URL from Lambda, upload the object directly to S3, and then save only the object metadata through the API. Realtime messages should use a service designed for managed persistent connections rather than a Socket.IO server running inside Lambda.')
add_table(doc, ['Component', 'Recommended service', 'Purpose'], [
    ('HTTP endpoints', 'API Gateway HTTP API + Lambda', 'Expose the existing REST-style routes without an always-on server'),
    ('Database', 'Supabase PostgreSQL', 'Avoid the fixed monthly cost of a separate RDS instance'),
    ('Object storage', 'Amazon S3', 'Store medical files and other large objects'),
    ('Realtime', 'API Gateway WebSocket API or IVS/managed realtime', 'Support bidirectional events and active-room features'),
    ('Background jobs', 'S3 event + SQS + worker', 'Process uploads asynchronously and retry failed jobs'),
    ('Monitoring', 'CloudWatch Logs and AWS Budgets', 'Track errors, duration, invocations, and spending'),
], widths=[1.45, 2.25, 3.1])

doc.add_heading('Expected monthly cost', level=1)
doc.add_paragraph('AWS Lambda charges for requests and execution duration. The Lambda free tier includes 1 million requests and 400,000 GB-seconds per month. API Gateway HTTP APIs have a 1 million request free tier for eligible accounts and are priced at approximately 1 USD per million requests after the free tier in the first pricing tier. Actual prices vary by region, account eligibility, data transfer, and tax.')
add_table(doc, ['Cost item', 'Likely starting cost', 'Main risk'], [
    ('Lambda', 'Often under a few USD for light-to-moderate API traffic', 'Long-running functions or high memory settings'),
    ('API Gateway HTTP API', 'Often near zero at this scale; approximately 1 USD per million requests after free usage', 'Large payloads are metered in data increments'),
    ('S3', 'Usually cents to a few USD for modest storage and requests', 'Large downloads, retention, and versioning'),
    ('Supabase', 'Existing provider cost', 'Database plan or connection limits'),
    ('IVS', 'Separate variable cost', 'Viewer-hours, ingest, and delivery can exceed the budget'),
    ('Conversion worker', 'Separate variable cost', 'CPU time and large files'),
], widths=[1.55, 2.5, 2.75])
doc.add_paragraph('A realistic initial budget is Lambda plus API Gateway plus modest S3 usage, with AWS Budgets set at 10 USD and 20 USD. Treat IVS and conversion as separate cost centers so they cannot silently consume the API budget.')

doc.add_heading('Migration plan', level=1)
for text in [
    'Create a Lambda entry point and remove app.listen from the request path. Wrap the Express app with serverless-http for the first migration so existing routes can be tested with minimal changes.',
    'Create an API Gateway HTTP API and connect it to the Lambda. Configure CORS, request throttling, access logging, and a custom domain only after the basic endpoint is stable.',
    'Move all credentials to Lambda environment variables or AWS Secrets Manager. Reuse Supabase clients and PostgreSQL pools outside the handler to reduce connection churn on warm invocations.',
    'Replace large multipart uploads with an endpoint that returns an S3 presigned URL. The browser uploads directly to S3; Lambda receives only metadata and authorization requests.',
    'Choose a realtime replacement. API Gateway WebSocket API is the closest AWS-native option, but the client protocol and Socket.IO server behavior will need to change.',
    'Move volume conversion to an asynchronous workflow. S3 upload events can enqueue work, and a worker can process the file, write the output to S3, and update status in Supabase.',
    'Load-test with realistic concurrency and configure alarms for Lambda errors, throttles, duration, API Gateway 5xx responses, and S3/IVS usage.',
]: add_bullet(doc, text)

doc.add_heading('Important constraints', level=1)
add_bullet(doc, 'Lambda is not a persistent Node.js server. Socket.IO connections cannot be maintained in the usual way inside a Lambda function.')
add_bullet(doc, 'Lambda has a maximum execution timeout of 900 seconds. Any conversion that can exceed 15 minutes must run in a separate worker or be broken into resumable steps.')
add_bullet(doc, 'Three thousand registered users is not the same as three thousand concurrent users. If all 3,000 users are simultaneously active, test concurrency, database connection behavior, WebSocket limits, and IVS usage before committing to the budget.')
add_bullet(doc, 'Do not put Lambda in a private VPC unless required. A NAT Gateway can cost more than the entire target budget.')

doc.add_heading('Security actions before deployment', level=1)
doc.add_paragraph('The repository contains a plaintext Supabase database credential in Medica/v1/utils/conn.js. Rotate that credential immediately, remove it from the repository and its history where practical, and deploy only a new value through environment variables or Secrets Manager. Treat the Supabase service-role key as a server-only secret and never expose it to the browser.')
add_bullet(doc, 'Use least-privilege IAM for Lambda, S3, logging, and any conversion worker.')
add_bullet(doc, 'Keep the S3 bucket private and issue short-lived presigned URLs.')
add_bullet(doc, 'Enable S3 lifecycle rules for incomplete uploads and old versions where retention policy permits.')
add_bullet(doc, 'Set AWS Budgets and billing alerts before production traffic is enabled.')

doc.add_heading('Decision', level=1)
doc.add_paragraph('Proceed with Lambda for the stateless HTTP API, API Gateway HTTP API as the entry point, Supabase as the database, and S3 for objects. Plan a deliberate redesign for realtime connections and volume conversion. This is the most credible AWS approach for the requested budget, but the 20 USD ceiling must exclude or separately control IVS traffic and heavy file processing.')

doc.add_heading('Reference sources', level=1)
sources = [
    'AWS Lambda pricing: https://aws.amazon.com/lambda/pricing/',
    'Amazon API Gateway pricing: https://aws.amazon.com/api-gateway/pricing/',
    'API Gateway HTTP APIs: https://docs.aws.amazon.com/apigateway/latest/developerguide/http-api.html',
    'API Gateway WebSocket APIs: https://docs.aws.amazon.com/apigateway/latest/developerguide/apigateway-websocket-api-overview.html',
    'Lambda timeout and limits: https://docs.aws.amazon.com/lambda/latest/dg/gettingstarted-limits.html',
    'Amazon S3 pricing: https://aws.amazon.com/s3/pricing/',
]
for s in sources: add_bullet(doc, s)

footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
fr = footer.add_run('AWS Lambda Scaling Recommendation')
fr.font.size = Pt(8); fr.font.color.rgb = RGBColor(100,100,100)

doc.core_properties.title = 'AWS Lambda Scaling Recommendation'
doc.core_properties.subject = 'Low-cost serverless architecture for the Project ANU application'
doc.core_properties.author = 'OpenAI'
doc.save(OUT)
print(OUT)
