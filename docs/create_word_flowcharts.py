from pathlib import Path
import json,re
from PIL import Image,ImageDraw,ImageFont
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
source=(ROOT/'ANU_SIMPLE_CLICK_FLOWCHARTS.html').read_text(encoding='utf8')
flows=json.loads(re.search(r'<script id="data" type="application/json">(.*?)</script>',source,re.S).group(1))
assets=ROOT.parent/'scratch'/'word-flowcharts'
assets.mkdir(parents=True,exist_ok=True)
out=ROOT/'ANU_Simple_Click_Flowcharts.docx'
font=ImageFont.truetype('C:/Windows/Fonts/arial.ttf',42)
bold=ImageFont.truetype('C:/Windows/Fonts/arialbd.ttf',38)

def wrap(text,width):
    rows=[];line=''
    for word in text.split():
        trial=(line+' '+word).strip()
        if font.getlength(trial)>width and line:rows.append(line);line=word
        else:line=trial
    if line:rows.append(line)
    return rows

def diagram(flow,index):
    layouts=[];height=12
    for s in flow['steps']:
        rows=wrap(s,1740);h=32+len(rows)*54
        layouts.append((rows,h));height+=h+42
    im=Image.new('RGB',(2000,height-20),'white');d=ImageDraw.Draw(im);y=12
    for j,(rows,h) in enumerate(layouts):
        d.rounded_rectangle((12,y,1988,y+h),radius=14,fill='#f3f7f9',outline='#b4c7d2',width=3)
        d.ellipse((36,y+20,96,y+80),fill='#224e64')
        n=str(j+1);d.text((66-bold.getlength(n)/2,y+26),n,font=bold,fill='white')
        for k,line in enumerate(rows):d.text((126,y+13+k*54),line,font=font,fill='#142f3d')
        y+=h
        if j<len(layouts)-1:
            d.line((1000,y+5,1000,y+29),fill='#65899b',width=4)
            d.polygon([(990,y+23),(1010,y+23),(1000,y+36)],fill='#65899b')
        y+=42
    p=assets/f'flow-{index:02}.png';im.save(p);return p

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(8.27);sec.page_height=Inches(11.69)
sec.top_margin=sec.bottom_margin=Inches(.6);sec.left_margin=sec.right_margin=Inches(.65)
for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
    st=doc.styles[name];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
    for border in list(st.element.xpath('.//w:pBdr')):
        border.getparent().remove(border)
doc.styles['Normal'].font.size=Pt(10.5)
doc.styles['Normal'].paragraph_format.space_after=Pt(6)
doc.styles['Normal'].paragraph_format.line_spacing=1.06
doc.styles['Title'].font.size=Pt(28)
doc.styles['Heading 1'].font.size=Pt(19)
doc.styles['Heading 2'].font.size=Pt(12)
doc.core_properties.title='ANU simple click flowcharts'
doc.core_properties.subject='User actions and cloud cost touchpoints'
doc.core_properties.author='Project ANU'
doc.add_paragraph('ANU simple click flowcharts','Title')
doc.add_paragraph('User actions and cloud cost touchpoints','Subtitle')
doc.add_paragraph('30 October 2026')
doc.add_paragraph('Use this guide to follow the application one task at a time. Each numbered box shows what to click and what appears next. Alternative choices explain errors, cancellation and ways to go back. Cloud notes explain where storing, downloading or sending information can add cost.')
doc.add_paragraph('Available screens and buttons depend on your role. Incomplete features and headset steps that require confirmation are marked in the relevant tasks.')
doc.add_heading('Find a section',1)
groups=list(dict.fromkeys(f['group'] for f in flows))
for g in groups:
    inds=[i+1 for i,f in enumerate(flows) if f['group']==g]
    doc.add_paragraph(f'{g}    Tasks {inds[0]} to {inds[-1]}')
doc.add_paragraph('Cloud costs are usage based. A click does not have a single fixed price. Charges depend on file size, retention, downloads, server work, live connection time and the provider plan.')

for i,f in enumerate(flows,1):
    doc.add_page_break()
    p=doc.add_paragraph(f['group']);p.paragraph_format.space_after=Pt(3)
    p.runs[0].bold=True;p.runs[0].font.size=Pt(10)
    title=re.sub(r'[^\w\s]',' ',f['title']).replace('_',' ')
    title=re.sub(r'\s+',' ',title).strip()
    doc.add_heading(f'{i} {title}',1)
    p=doc.add_paragraph();p.paragraph_format.space_after=Pt(5)
    picture=p.add_run().add_picture(str(diagram(f,i)),width=Inches(6.95))
    picture._inline.docPr.set('descr','Click sequence for '+f['title']+'. '+' '.join(f'{j+1}. {s}' for j,s in enumerate(f['steps'])))
    if f['branches']:
        doc.add_heading('Other choices',2)
        for step,label,result in f['branches']:
            p=doc.add_paragraph();r=p.add_run(f'After step {step+1}  {label}. ');r.bold=True;p.add_run(result)
    if f['cost']:
        p=doc.add_paragraph();p.add_run('Cloud cost here. ').bold=True;p.add_run(f['cost'])
    if f['note']:
        p=doc.add_paragraph();p.add_run('Current project note. ').bold=True;p.add_run(f['note'])

doc.add_page_break();doc.add_heading('Cloud pricing reference',1)
doc.add_paragraph('Amazon IVS pricing: https://aws.amazon.com/ivs/pricing/')
doc.add_paragraph('The live-video example uses the India first-tier rate of USD 0.092 per connected participant-hour. The full bill also depends on replication, other cloud resources, applicable tiers, allowances, discounts and taxes.')
doc.add_paragraph('Storage charges depend on average retained file size and storage duration. Download charges depend on file requests and delivered data. Conversion charges depend on computer resources and running time. Database, server and notification charges depend on the selected provider plans and actual usage.')
footer=sec.footer.paragraphs[0];footer.alignment=2
r=footer.add_run();field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');r._r.addnext(field)
doc.save(out)
print(out)
