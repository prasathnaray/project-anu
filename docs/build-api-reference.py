from pathlib import Path
from collections import defaultdict
import json,re
from docx import Document
from docx.shared import Inches,Pt,RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT=Path(__file__).resolve().parent
inventory=json.loads((ROOT/'ANU_API_INVENTORY.json').read_text())
entries=inventory['endpoints']

def words(s):
 s=re.sub(r'([a-z0-9])([A-Z])',r'\1 \2',s)
 return re.sub(r'\s+',' ',s.replace('_',' ')).strip()

def group(e):
 f=Path(e['routeFile']).stem.lower();p=e['endpoint'].lower()
 if p in ['/health','/test']:return 'Service checks'
 if f in ['auth','forgotproute','accountverificationroute','refreshtokenrouter','sessionroute']:return 'Login account verification and device sessions'
 if f in ['streamingsessionroute','tokenivsroute','individualivsroute','activepeopleroomroute']:return 'Headset streaming and live viewing'
 if f in ['gettraineesroute','createtraineeroute','disabletraineeroute','deletetraineeroute','updatetraineeroute','induuidroute','getinstructorroute','deleteinstructorroute','updateinstructorroute','getinstructorasperbatch','getadminsroute']:return 'Trainees instructors and user accounts'
 if 'targeted' in f or f=='getindividualtlroute':return 'Targeted learning'
 if 'batch' in f:return 'Batches'
 if f in ['superadminroute','institutionroute','createscancentersroute','createscancenterroute','getscancenterroute']:return 'Institutions and platform administration'
 if f=='contentaccessroute':return 'Course access publication assignments and migration'
 if f=='mrcontentroute':return 'Headset content packages and recording validation'
 if f=='clusterroute':return 'Volume clusters and cluster recordings'
 if f=='coursemappingroute':return 'Volume course mappings'
 if any(x in f for x in ['volume','svupload','shadowrecording']):return 'Volumes placement reference recordings and downloads'
 if f in ['iivrroutes','iivrstarttestroute','iivrendtestroute','practiceroute','practestroute','getsubmissionsroute','getpracticeresults']:return 'Image interpretation practice tests and submissions'
 if f in ['questionsroute','mindsparkroute','lrobroute','challengeroute']:return 'MindSpark OB Booster and challenges'
 if f in ['notificationroutes','queriesroute']:return 'Queries and notifications'
 if f=='activitytrackerroute':return 'Activity tracking'
 if f=='reattroute':return 'Reattempt settings'
 if f in ['profileroute','dashboardroute','userstatusroute','getgendersroutes','getinstructordataanalysisroute','interactionsattemptstats','activitylastscoresroute','skillcompetencyroute','performancemetricsroute']:return 'Profiles dashboards and analytics'
 return 'Curriculum certifications modules resources and progress'

special={
 '/health':'Check whether the service responds.', '/test':'Return authenticated test-route and user information.',
 '/api/v1/login':'Sign in and receive tokens plus account and device context.',
 '/api/v1/forgot-password':'Record a password-recovery request and attempt an email.',
 '/api/v1/account-verification/request':'Request an activation or reactivation phone OTP.',
 '/api/v1/account-verification/verify':'Verify an account activation or reactivation OTP.',
 '/api/v1/refresh-token':'Issue a new access token using the refresh-token cookie.',
 '/api/v1/tokenn':'Create a headset publisher session and IVS participant token.',
 '/api/v1/streaming/self-viewer-session':'Create a trainee self-preview viewer session.',
 '/api/v1/streaming/viewer-session':'Create an institution viewer session.',
 '/api/v1/streaming/publisher-session/:sessionId/activate':'Activate a published headset session and its replication.',
 '/api/v1/streaming/publisher-session/:sessionId':'Stop the headset publisher session.',
 '/api/v1/submit-msob':'Submit a MindSpark or OB Booster activity result.',
 '/api/v1/submit-ii':'Submit an image-interpretation, annotation or measurement result.',
 '/api/v1/start-ii-prac':'Start an image-interpretation practice or test session.',
 '/api/v1/end-ii-prac':'End an image-interpretation session using test_id.',
 '/api/v1/resource-completion':'Mark resource completion through query parameters.',
 '/api/v1/get-vr-data/:data_check':'Return authenticated user context and the supplied data_check value.',
 '/api/v1/send':'Create an in-app notification.',
 '/api/v1/create-module':'Create a chapter using chapter_name and course_id.',
 '/api/v1/create-new-module':'Create a module underneath a chapter.',
 '/api/v1/mind-spark-options':'Save a legacy MindSpark option response.',
 '/api/v1/mr/workspace':'Retrieve the requester’s permitted headset content workspace.',
 '/api/v1/mr/courses':'Retrieve courses available to the headset content workspace.',
 '/api/v1/mr/courses/:courseId/package':'Retrieve the authorized course content package.',
 '/api/v1/trainee/update':'Update trainee batch association using query parameters.',
 '/api/v1/mark-testscores-re':'Update reattempt scores using query parameters.',
}
def purpose(e):
 if e['endpoint'] in special:return special[e['endpoint']]
 h=e['handlerExpression']
 if re.fullmatch(r'[\w.]+',h):
  h=h.split('.')[-1]
  h=re.sub(r'(Controller|controller|Cont|Request)$','',h)
  h=re.sub(r'(?<=[a-z])C$','',h)
  s=words(h)
  if s.lower() not in ['list','create','get stats','get global statistics','get activity summary','get recent activities']:
   return s[0].upper()+s[1:]+'.'
 p=e['endpoint'].replace('/api/v1/','').replace('/',' ')
 p=re.sub(r':\w+','selected item',p)
 verb={'GET':'Retrieve','POST':'Create or execute','PUT':'Update','PATCH':'Update','DELETE':'Delete'}.get(e['method'],'Access')
 return verb+' '+p+'.'

for e in entries:
 e['params']=re.findall(r':([A-Za-z_]\w*)',e['endpoint'])
 if e['service']=='Main LMS' and e['endpoint']=='/api/v1/reatt-data' and e['method']=='GET':
  e['query']=['uploaded_by','certificate_id','course_id','unit_name','resource_type','resource_id','max_reattempt_count']
 if e['endpoint'].endswith(('/shadow-recordings','/step-recordings')) and '/clusters/' in e['endpoint']:
  e['query']=[x for x in e['query'] if x!='recording_type']
 e['group']=group(e);e['purpose']=purpose(e)
 if e['endpoint']=='/api/v1/submit-msob' and e['service']=='Main LMS':
  e['body']=sorted(set(e['body']+['questionNo','optionChosen','isCorrect','payload','timeTaken','hasTakenClue','totalTimeTaken']))
 e['body']=[x for x in e['body'] if x not in ('undefined',)]

notes={
 '/api/v1/forgot-password':'Current SMTP message is placeholder text; this route does not complete a secure password reset.',
 '/api/v1/resource-completion':'This GET route changes completion state. Query keys are shown exactly as implemented.',
 '/api/v1/trainee/update':'This PUT route reads user_id and batch_id from the query, not a JSON body.',
 '/api/v1/mark-testscores-re':'This PUT route reads the score fields from query parameters.',
 '/api/v1/submit-msob':'Fields depend on resourceType. See the activity payload guide below.',
 '/api/v1/tokenn':'Main LMS contract ignores client-provided stage/user/capability choices. Medica is a separate implementation.',
 '/api/v1/submit-ii':'Required fields depend on questionType; see the submission guide below.',
 '/api/v1/send':'No shared bearer middleware is mounted on this route in app.js.',
 '/api/v1/activity-logs/track':'softAuth decodes a supplied token; it does not verify its signature in this route middleware.',
}

def details(e):
 parts=[]
 for key,label in [('params','Path'),('query','Query'),('body','Body'),('upload','Files')]:
  if e[key]:parts.append(label+': '+', '.join(e[key]))
 if not parts:parts.append('No named request fields identified in the traced handler.')
 if e['upload']:parts.append('Send multipart/form-data.')
 if e['status']:parts.append('Observed status codes: '+', '.join(str(x) for x in e['status']))
 else:parts.append('Status / payload follows handler and shared error handling.')
 parts.append('Source: '+e['routeFile'].split('/routes/')[-1]+':'+str(e['routeLine']))
 return '\n'.join(parts)

doc=Document();sec=doc.sections[0]
sec.page_width=Inches(11.69);sec.page_height=Inches(8.27)
sec.top_margin=sec.bottom_margin=Inches(.55);sec.left_margin=sec.right_margin=Inches(.6)
for n in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
 st=doc.styles[n];st.font.name='Arial';st.font.color.rgb=RGBColor(0,0,0)
 for el in list(st.element.xpath('.//w:pBdr')):el.getparent().remove(el)
doc.styles['Normal'].font.size=Pt(10)
doc.styles['Normal'].paragraph_format.space_after=Pt(5)
doc.styles['Title'].font.size=Pt(26)
doc.styles['Heading 1'].font.size=Pt(17)
doc.styles['Heading 2'].font.size=Pt(12)
doc.core_properties.title='ANU API endpoint reference'
doc.core_properties.subject='Registered routes methods access and request fields'
doc.core_properties.author='Project ANU'
doc.add_paragraph('ANU API endpoint reference','Title')
doc.add_paragraph('Main LMS and Medica service routes','Subtitle')
doc.add_paragraph('30 October 2026')
doc.add_paragraph('This reference lists the registered HTTP APIs in the project, grouped by feature. Use the method and full endpoint together to call an API. The request columns preserve the field names used in the code, including older spellings and aliases.')
doc.add_paragraph('Main LMS: 194 method and endpoint combinations from server/v1/app.js. Medica companion: 49 combinations from Medica/v1/app.js. They are separate services even when endpoint paths match. Aliases are counted separately. The empty server/v2/login files define no API. One additional route file is not mounted and is listed separately.')
doc.add_heading('How to call an endpoint',1)
doc.add_paragraph('Full URL = service base URL + endpoint. For local main LMS development, the default base is http://localhost:4004, unless PORT is changed. Use the configured deployment host in production. Medica also declares port 4004 in its own app, so give the services distinct hosts or ports if both run together.')
doc.add_paragraph('Example: GET http://localhost:4004/api/v1/get-trainees?page=1&limit=10')
doc.add_paragraph('Bearer endpoints require Authorization: Bearer <access_token>. JSON requests use Content-Type: application/json. File requests use multipart/form-data; let the client create its boundary. A :parameter in a path must be replaced with the real value and URL-encoded as necessary.')
doc.add_heading('Reading the reference',1)
doc.add_paragraph('Bearer means the route is behind the shared authentication middleware. Role, institution, ownership and assignment checks can still restrict access. Roles in this project are 99 super admin, 101 institution admin, 102 instructor and 103 trainee. A listed field is read by the code; it is not automatically required. Conditional or nested payload rules are explained for key submissions below.')
doc.add_paragraph('No shared bearer middleware describes route registration, not a guarantee that the controller permits every request. Optional decoded token identifies the activity tracker’s soft-auth path. Refresh-token cookie identifies the refresh route. Status codes shown are literal codes found in the traced handler path, not an exhaustive response contract; middleware, validation and database errors may add other responses.')
doc.add_paragraph('The endpoint list follows source registration and does not certify deployment or successful execution. Response structures vary across older and newer routes. No universal response envelope should be assumed.')

def table(headers,widths):
 t=doc.add_table(rows=1,cols=len(headers));t.autofit=False
 for col,w in zip(t.columns,widths):col.width=Inches(w)
 for c,w,h in zip(t.rows[0].cells,widths,headers):c.width=Inches(w);c.text=h
 repeat=OxmlElement('w:tblHeader');t.rows[0]._tr.get_or_add_trPr().append(repeat)
 borders=OxmlElement('w:tblBorders')
 for k in ['top','left','bottom','right','insideH','insideV']:
  v=OxmlElement('w:'+k);v.set(qn('w:val'),'single');v.set(qn('w:sz'),'4');v.set(qn('w:color'),'D9D9D9');borders.append(v)
 t._tbl.tblPr.append(borders)
 margins=OxmlElement('w:tblCellMar')
 for k in ['top','left','bottom','right']:
  v=OxmlElement('w:'+k);v.set(qn('w:w'),'85');v.set(qn('w:type'),'dxa');margins.append(v)
 t._tbl.tblPr.append(margins)
 for c in t.rows[0].cells:
  shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'244759');c._tc.get_or_add_tcPr().append(shade)
  for r in c.paragraphs[0].runs:r.bold=True;r.font.color.rgb=RGBColor(255,255,255);r.font.size=Pt(9.5)
 return t

def bodyrow(t,values,idx):
 row=t.add_row();trpr=row._tr.get_or_add_trPr();trpr.append(OxmlElement('w:cantSplit'))
 for c,value in zip(row.cells,values):
  c.text=value;c.vertical_alignment=1
  if idx%2==0:
   shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'F1F5F7');c._tc.get_or_add_tcPr().append(shade)
  for p in c.paragraphs:
   p.paragraph_format.space_after=Pt(3);p.paragraph_format.line_spacing=1.03
   for r in p.runs:r.font.size=Pt(9)
 return row

group_order=list(dict.fromkeys(e['group'] for e in entries if e['service']=='Main LMS'))
for service in ['Main LMS','Medica companion']:
 groups=defaultdict(list)
 for e in entries:
  if e['service']==service:groups[e['group']].append(e)
 for g in group_order+[k for k in groups if k not in group_order]:
  if not groups[g]:continue
  doc.add_heading(service+' '+g.lower(),1).paragraph_format.page_break_before=True
  doc.add_paragraph(str(len(groups[g]))+' registered method and endpoint combinations. Source paths below are relative to '+('server/v1/routes/' if service=='Main LMS' else 'Medica/v1/routes/')+' unless an app.js path is shown.')
  t=table(['Method','Full endpoint','Purpose and access','Request details and observed status codes'],[.62,2.92,2.15,4.79])
  for idx,e in enumerate(groups[g]):
   access={'Bearer token':'Bearer','No shared bearer middleware':'No shared bearer middleware','Optional decoded token':'Optional decoded token','Refresh-token cookie / handler check':'Refresh-token cookie'}[e['auth']]
   bodyrow(t,[e['method'],e['endpoint'],e['purpose']+'\nAccess: '+access,details(e)],idx)
  relevant=list(dict.fromkeys(notes[e['endpoint']] for e in groups[g] if e['endpoint'] in notes))
  for n in relevant:doc.add_paragraph(n)

doc.add_heading('Request examples for main LMS',1).paragraph_format.page_break_before=True
doc.add_paragraph('Replace placeholders with real authorized values. These examples illustrate the source contract; they have not been submitted to the running service.')
examples=[
 ('Sign in','POST /api/v1/login',{'user_mail':'trainee@example.com','user_password':'<password>'}),
 ('Request verification OTP','POST /api/v1/account-verification/request',{'email':'trainee@example.com','purpose':'activation'}),
 ('Verify an account','POST /api/v1/account-verification/verify',{'email':'trainee@example.com','purpose':'activation','otp':'123456'}),
 ('Start image interpretation practice','POST /api/v1/start-ii-prac',{'resource_id':'<resource-id>'}),
 ('End image interpretation practice','PUT /api/v1/end-ii-prac',{'test_id':'<test-id-returned-by-the-session-flow>'}),
 ('Create a query','POST /api/v1/queries',{'subject':'Learning activity question','message':'Please clarify the expected plane.'}),
 ('Create a modern course draft','POST /api/v1/courses',{'name':'Example course','courseKind':'institution','curriculumId':'<curriculum-id>'}),
 ('Set course assignments','PUT /api/v1/courses/:courseId/assignments',{'batchIds':['<batch-id>'],'assignedTraineeIds':[],'excludedTraineeIds':[]}),
]
for title,endpoint,payload in examples:
 doc.add_heading(title,2);doc.add_paragraph(endpoint)
 p=doc.add_paragraph(json.dumps(payload,indent=2));p.paragraph_format.space_after=Pt(10)
 for r in p.runs:r.font.name='Consolas';r.font.size=Pt(9)

doc.add_heading('Conditional submission payloads',1).paragraph_format.page_break_before=True
doc.add_heading('Image interpretation submission',2)
doc.add_paragraph('POST /api/v1/submit-ii accepts questionType, questionNo, session_id and resource_id as the common required fields. isCorrect is required except for measurement. Upload userImage using multipart/form-data for type2, annotation1, annotation2 and measurement.')
t=table(['questionType','Additional requirements'],[2,8.48])
for i,row in enumerate([
 ('type1','optionChosen is checked for question numbers below 6. No userImage is required by this branch.'),
 ('type2','userImage and isCorrect.'),
 ('annotation1 or annotation2','userImage, isCorrect, correctLabelCount, wrongLabelCount and unusedLabelCount.'),
 ('measurement','userImage, partial, value, interpretation and caliperPlacementInterpretation. partial must be 0, 0.5 or 1.'),
]):bodyrow(t,row,i)
doc.add_heading('MindSpark and OB Booster submission',2)
doc.add_paragraph('POST /api/v1/submit-msob requires resourceType and resourceId. sessionId can be supplied; the controller generates a UUID when absent. Resource types are normalized to uppercase.')
t=table(['resourceType','Activity specific fields'],[2,8.48])
for i,row in enumerate([
 ('TYPE1','questionNo, optionChosen, isCorrect'),('MATCHING','questionNo, payload, isCorrect'),('WORDSEARCH','questionNo, timeTaken, hasTakenClue'),('CROSSWORD','totalTimeTaken'),('PROBEMOVEMENTS','totalTimeTaken')
]):bodyrow(t,row,i)
doc.add_heading('Challenge submission',2)
doc.add_paragraph('POST /api/v1/challenges/submit validates resource_id, session_id, question_number, choose_option and isCorrect. question_number must be numeric. choose_option can be a single value or an array and cannot be empty after normalization. isCorrect must parse as true or false.')

doc.add_heading('File uploads and headset communication',1).paragraph_format.page_break_before=True
doc.add_heading('File request rules',2)
doc.add_paragraph('The file-field names in the endpoint tables are exact multipart keys. Ordinary form fields accompany them. The route may accept alternative file key names; aliases are not a requirement to send the same file multiple times.')
doc.add_paragraph('Volume uploads use file. Placement uses placed_file. Volume-reference recording upload uses recording_file, audio_file, images and manifest_file. Cluster recording routes accept additional aliases shown in their table. Question image assets use image. Image-interpretation submissions use userImage. Practice-test submissions accept the numbered userImage and expertImage keys and plural aliases shown in the table.')
doc.add_paragraph('The volume upload limit is environment-controlled. The tracked Lambda template sets MAX_VOLUME_UPLOAD_SIZE_MB to 8, while source defaults differ. Recording routes specify 50 MB per file and a total file-count limit of 61. Question assets default to 10 MB through MAX_MINDSPARK_ASSET_SIZE_MB. Deployed proxy limits can be lower; use the actual environment when choosing an upload size.')
doc.add_heading('Live headset sequence',2)
for s in ['1. Headset calls POST /api/v1/tokenn with its LMS bearer token.','2. The publisher token is returned at result.data.token and the application session ID at result.sessionId.','3. Headset joins IVS and publishes video and microphone.','4. Activation is retried in the background. The headset can call POST /api/v1/streaming/publisher-session/:sessionId/activate; a 409 means publishing is not yet visible and needs bounded retry.','5. Trainee browser calls POST /api/v1/streaming/self-viewer-session. Institution-admin browser calls POST /api/v1/streaming/viewer-session.','6. Headset calls DELETE /api/v1/streaming/publisher-session/:sessionId when publishing stops.']:
 doc.add_paragraph(s)
doc.add_paragraph('IVS media transport is separate from the REST API endpoints. A storage download endpoint generally prepares an authorized URL; the subsequent file transfer can occur directly from object storage.')
doc.add_heading('Socket events',2)
doc.add_paragraph('The local main LMS also initializes Socket.IO. Client subscriptions are subscribe:superadmin and subscribe:user. Server events include activity:new, stats:update and session:revoked. These are event names, not standalone REST endpoints. The Lambda HTTP deployment does not by itself establish a persistent Socket.IO host.')

doc.add_heading('Routes not registered and reference sources',1).paragraph_format.page_break_before=True
for e in inventory['unmounted']:
 doc.add_paragraph(e['service']+' defines '+e['method']+' '+e['relativeRoute']+' in '+e['file']+':'+str(e['line'])+', but its router is not mounted in the service app.js. No public full endpoint is assigned here.')
doc.add_paragraph('server/v2/login/endpoint.js, authRoute.js, authController.js and authModel.js are empty. They contribute no endpoints. The volume converter is a Batch worker in the reviewed deployment and is not an additional HTTP service.')
doc.add_paragraph('Primary sources are the two app.js files, their imported route modules, and the associated controllers, models and authentication middleware. The endpoint registration list excludes commented-out code and unmounted route files. The route inventory and source line numbers were generated from JavaScript syntax trees, then supplemented for conditional payloads.')
doc.add_paragraph('Useful source documents: server/v1/STREAMING.md, server/v1/template.yaml, APIs/Submit_MSOB_API_Product_Guide.md, APIs/Volume_APIs.md and APIs/Challenges_API.md. Endpoint spellings are taken from the current route code when older documents differ.')
footer=sec.footer.paragraphs[0];footer.alignment=2
footer.add_run('ANU API reference  |  ')
field=OxmlElement('w:fldSimple');field.set(qn('w:instr'),'PAGE');footer._p.append(field)
out=ROOT/'ANU_API_Endpoint_Reference.docx';doc.save(out)

md='# ANU API endpoint reference\n\n194 main LMS routes and 49 Medica companion routes. Fields are observed inputs, not a full validation schema.\n'
for service in ['Main LMS','Medica companion']:
 md+='\n## '+service+'\n'
 for e in entries:
  if e['service']!=service:continue
  md+=f'\n### {e["method"]} {e["endpoint"]}\n\n{e["purpose"]}\n\nAccess: {e["auth"]}\n\n'+details(e)+'\n'
(ROOT/'ANU_API_Endpoint_Reference.md').write_text(md,encoding='utf8')
print(out)
print('Endpoint rows:',len(entries),'groups:',len(group_order))
