from pathlib import Path
import pypdfium2 as pdfium
from pypdf import PdfReader
from docx import Document
from docx.shared import Pt
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

root=Path(__file__).resolve().parent
source=Path('C:/Users/igrs/Downloads/ANU — Simple click flowcharts.pdf')
work=root.parent/'scratch'/'exact-pdf-word'
work.mkdir(parents=True,exist_ok=True)
pdf=pdfium.PdfDocument(str(source));reader=PdfReader(source)
sizes={(round(float(p.mediabox.width),2),round(float(p.mediabox.height),2)) for p in reader.pages}
assert len(sizes)==1,sizes
w,h=next(iter(sizes))
doc=Document();sec=doc.sections[0]
sec.page_width=Pt(w);sec.page_height=Pt(h)
sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Pt(0)
sec.header_distance=sec.footer_distance=Pt(0)
normal=doc.styles['Normal'];normal.font.size=Pt(1)
normal.paragraph_format.space_before=normal.paragraph_format.space_after=Pt(0)
doc.core_properties.title='ANU simple click flowcharts'
doc.core_properties.subject='Exact page layout from the supplied PDF'
for i in range(len(pdf)):
    image=work/f'pdf-page-{i+1:02}.png'
    pdf[i].render(scale=2.5).to_pil().convert('RGB').save(image)
    p=doc.paragraphs[0] if i==0 and doc.paragraphs else doc.add_paragraph()
    p.paragraph_format.page_break_before=i>0
    shape=p.add_run().add_picture(str(image),width=Pt(w),height=Pt(h))
    inline=shape._inline
    inline.docPr.set('descr',reader.pages[i].extract_text() or f'Source PDF page {i+1}')
    anchor=OxmlElement('wp:anchor')
    for key,value in {'distT':'0','distB':'0','distL':'0','distR':'0','simplePos':'0','relativeHeight':'0','behindDoc':'1','locked':'0','layoutInCell':'1','allowOverlap':'1'}.items():anchor.set(key,value)
    pos=OxmlElement('wp:simplePos');pos.set('x','0');pos.set('y','0');anchor.append(pos)
    for tag in ['positionH','positionV']:
        pos=OxmlElement('wp:'+tag);pos.set('relativeFrom','page')
        offset=OxmlElement('wp:posOffset');offset.text='0';pos.append(offset);anchor.append(pos)
    for child in list(inline):
        if child.tag==qn('wp:docPr'):anchor.append(OxmlElement('wp:wrapNone'))
        anchor.append(child)
    inline.getparent().replace(inline,anchor)
out=root/'ANU_Exact_PDF_Content.docx';doc.save(out)
print(f'Created {out} with {len(pdf)} source pages')
